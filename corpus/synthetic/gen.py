"""Generate a synthetic ~200-page docx (python-docx), Dutch-ish legal prose, headings, lists, tables."""
import random, sys
from docx import Document
from docx.shared import Pt

random.seed(1)
words = ("de werknemer werkgever arbeidsovereenkomst artikel bepaling salaris uur per week vakantie verlof "
         "partijen overeenkomst geheimhouding opzegging termijn maand jaar recht plicht toepassing conform "
         "wet regeling functie werkzaamheden locatie Utrecht vergoeding reiskosten pensioen ziekte melding").split()

def sent():
    n = random.randint(8, 22)
    s = " ".join(random.choice(words) for _ in range(n))
    return s[0].upper() + s[1:] + "."

pages = int(sys.argv[2]) if len(sys.argv) > 2 else 200
doc = Document()
art = 0
# ~ 450 words/page
total_words = pages * 450
w = 0
while w < total_words:
    art += 1
    doc.add_heading(f"Artikel {art} {random.choice(words).capitalize()}", level=1)
    for _ in range(random.randint(2, 5)):
        p = doc.add_paragraph()
        txt = " ".join(sent() for _ in range(random.randint(2, 5)))
        r = p.add_run(txt[: len(txt)//2]); r2 = p.add_run(txt[len(txt)//2:])
        if random.random() < 0.2: r2.bold = True
        w += len(txt.split())
    for _ in range(random.randint(0, 3)):
        p = doc.add_paragraph(sent(), style="List Number"); w += 10
    if art % 15 == 0:
        t = doc.add_table(rows=6, cols=4)
        for i, row in enumerate(t.rows):
            for j, c in enumerate(row.cells):
                c.text = f"{random.randint(1000,9999)}" if i else ["Schaal", "Min", "Max", "Periodiek"][j]
        w += 24
doc.save(sys.argv[1])
print("articles", art, "words", w)
