"""Head-to-head fidelity benchmark: Quillwork, Adeu, Docxodus (docx-scalpel) and python-docx on the
public Quillwork corpus (corpus/synthetic/*.docx + corpus/fixtures/*.docx, 19 files).

    python bench/fidelity_compare.py [--json OUT.json] [--no-validate] [FILE.docx ...]
    python bench/compare_report.py          # per-file tables for results-compare.md

It reuses bench/fidelity.py (metrics: noop(), edit_locality(), paragraphs(), and the
target-word picker pick_word()), so every engine is measured the same way and edits the same word.

Per engine and file:

no-op     open, save, change nothing. file byte-identical / zip entries with identical stored bytes
          / parts with identical decompressed content.
edit      replace the picked word with "Edited", in the engine's tracked-change mode where it has one:
            Quillwork  Document.apply(replace_text, anchor = first find() hit)      -> tracked
            Adeu       RedlineEngine.process_batch([ModifyText(regex r"\\bword\\b",
                       match_mode="first")])                                          -> tracked
            Docxodus   open_session(tracked_changes=RENDER_INLINE); grep(r"\\bword\\b",
                       max_results=1) -> replace_match()                              -> tracked
            python-docx  overwrite runs[0].text of the first paragraph whose first run has the word
          other entries changed / outermost w:p paragraphs changed (fidelity.edit_locality), plus:
          target_ok   the picked paragraph's accepted text (tracked changes accepted) equals its
                      original accepted text with the word replaced once, and no other paragraph's
                      accepted text changed. (Markup-only changes elsewhere show up in
                      "paragraphs changed", not here.)
          kind        tracked (new w:del holding the word + new w:ins holding "Edited"),
                      overwrite (no new revision markup) or other.
valid     every output: opens as a zip, every XML part parses, python-docx can load it, and
          Open XML SDK validation (via Docxodus verify_deliverable, with the input as baseline so only
          errors the engine introduced count; categories package/openXml/relationship/structure).
canonical the same no-op and edit comparisons on a canonical form (C14N: inclusive per XML part,
          exclusive per outermost w:p; plus indentation-only whitespace outside w:t-like leaves dropped,
          and [Content_Types].xml / .rels children sorted), so serialisation churn (XML declaration,
          attribute order, `<x/>` vs `<x />` vs `<x></x>`, pretty-printing, entry order) is separated
          from content change.
tokens    o200k_base tokens of each engine's full read view: Quillwork read(budget=1e9), Adeu
          extract_text_from_stream (CriticMarkup), Docxodus project().markdown, python-docx paragraph
          text, raw word/document.xml.
"""

from __future__ import annotations

import argparse
import base64
import html
import io
import json
import logging
import platform
import re
import sys
import time
import traceback
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import Any

import structlog

structlog.configure(
    wrapper_class=structlog.make_filtering_bound_logger(logging.WARNING)
)

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))  # for bench/fidelity.py

import adeu
import docx
import quillwork
import tiktoken
from adeu import ModifyText, RedlineEngine, extract_text_from_stream
from docx_scalpel import open_session
from docx_scalpel._transport import call as scalpel_call
from docx_scalpel.enums import TrackedChangeMode
from docx_scalpel.types import DocxSessionSettings
from fidelity import ENC, FULL, MAIN, edit_locality, noop, paragraphs, parts, pick_word
from lxml import etree

ENGINES = ["quillwork", "adeu", "docxodus", "python-docx"]
REPLACEMENT = "Edited"
AUTHOR = "bench"
VALIDITY_CATEGORIES = {"package", "openXml", "relationship", "structure"}


# --------------------------------------------------------------------------- engines


def q_noop(data: bytes) -> bytes:
    return quillwork.Document.open(data).save()


def q_edit(data: bytes, word: str) -> bytes:
    q = quillwork.Document.open(data)
    q.set_now(1_767_225_600)
    hit = re.search(r"^\[([0-9a-z.]+)", q.find(word), re.MULTILINE)
    if not hit:
        raise RuntimeError("find() returned no hit")
    q.apply(
        {
            "base_version": q.version,
            "ops": [
                {
                    "op": "replace_text",
                    "anchor": hit.group(1),
                    "find": word,
                    "replace": REPLACEMENT,
                }
            ],
        }
    )
    return q.save()


def a_noop(data: bytes) -> bytes:
    return RedlineEngine(io.BytesIO(data), author=AUTHOR).save_to_stream().getvalue()


def a_edit(data: bytes, word: str) -> bytes:
    e = RedlineEngine(io.BytesIO(data), author=AUTHOR, timestamp="2026-01-01T00:00:00Z")
    st = e.process_batch(
        [
            ModifyText(
                target_text=rf"\b{word}\b",
                new_text=REPLACEMENT,
                match_mode="first",
                regex=True,
            )
        ]
    )
    if st.get("edits_applied") != 1:
        raise RuntimeError(f"adeu did not apply the edit: {json.dumps(st)[:500]}")
    return e.save_to_stream().getvalue()


def d_noop(data: bytes) -> bytes:
    with open_session(data) as s:
        return s.save()


def d_edit(data: bytes, word: str) -> bytes:
    settings = DocxSessionSettings(
        tracked_changes=TrackedChangeMode.RENDER_INLINE, revision_author=AUTHOR
    )
    with open_session(data, settings) as s:
        ms = s.grep(rf"\b{word}\b", max_results=1)
        if not ms:
            raise RuntimeError("grep found no match")
        r = s.replace_match(ms[0], REPLACEMENT)
        if not r.success:
            raise RuntimeError(f"replace_match failed: {r.error}")
        return s.save()


def p_noop(data: bytes) -> bytes:
    buf = io.BytesIO()
    docx.Document(io.BytesIO(data)).save(buf)
    return buf.getvalue()


def p_edit(data: bytes, word: str) -> bytes:
    d = docx.Document(io.BytesIO(data))
    for p in d.paragraphs:
        if p.runs and re.search(rf"\b{word}\b", p.runs[0].text):
            p.runs[0].text = re.sub(
                rf"\b{word}\b", REPLACEMENT, p.runs[0].text, count=1
            )
            break
    else:
        raise RuntimeError("no paragraph with the word in its first run")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


NOOP: dict[str, Callable[[bytes], bytes]] = {
    "quillwork": q_noop,
    "adeu": a_noop,
    "docxodus": d_noop,
    "python-docx": p_noop,
}
EDIT: dict[str, Callable[[bytes, str], bytes]] = {
    "quillwork": q_edit,
    "adeu": a_edit,
    "docxodus": d_edit,
    "python-docx": p_edit,
}
EDIT_MODE = {
    "quillwork": "tracked (replace_text)",
    "adeu": "tracked (ModifyText)",
    "docxodus": "tracked (RENDER_INLINE)",
    "python-docx": "overwrite (no tracked-change API)",
}


def read_view(engine: str, data: bytes) -> str:
    if engine == "quillwork":
        return quillwork.Document.open(data).read(budget=FULL)
    if engine == "adeu":
        return extract_text_from_stream(io.BytesIO(data))
    if engine == "docxodus":
        with open_session(data) as s:
            return s.project().markdown
    d = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in d.paragraphs)


# --------------------------------------------------------------------------- target + checks

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def target_index(path: Path, word: str) -> int:
    """Index (among outermost w:p of document.xml, the unit fidelity.paragraphs() counts) of the body
    paragraph pick_word() chose: the first top-level body paragraph whose first run holds the word."""
    d = docx.Document(str(path))
    outer = [
        p
        for p in d.element.body.iter(f"{W}p")
        if not any(a.tag == f"{W}p" for a in p.iterancestors())
    ]
    for i, block in enumerate(outer):
        if block.getparent() is d.element.body:
            runs = docx.text.paragraph.Paragraph(block, d).runs
            if runs and re.search(rf"\b{word}\b", runs[0].text):
                return i
    raise RuntimeError("picked paragraph not found")


DEL = re.compile(rb"<w:(del|moveFrom)(?=[\s>])[^>]*?(?<!/)>.*?</w:\1>", re.DOTALL)
T = re.compile(rb"<w:t(?:\s[^>]*)?>(.*?)</w:t>", re.DOTALL)


def accepted_text(p: bytes) -> str:
    """Paragraph text with all tracked changes accepted (w:del/w:moveFrom dropped, w:ins kept)."""
    return html.unescape(b"".join(T.findall(DEL.sub(b"", p))).decode("utf-8"))


def count(tag: str, p: bytes) -> int:
    return len(re.findall(rb"<w:" + tag.encode() + rb"(?=[\s>])", p))


def classify_edit(before: bytes, after: bytes, idx: int, word: str) -> dict[str, Any]:
    old_ps, new_ps = paragraphs(parts(before)[MAIN]), paragraphs(parts(after)[MAIN])
    if len(old_ps) != len(new_ps) or idx >= len(new_ps):
        return {"target_ok": False, "kind": "other", "why": "paragraph count changed"}
    old, new = old_ps[idx], new_ps[idx]
    expected = re.sub(rf"\b{word}\b", REPLACEMENT, accepted_text(old), count=1)
    text_ok = accepted_text(new) == expected
    new_del = count("del", new) - count("del", old)
    new_ins = count("ins", new) - count("ins", old)
    deltexts = "".join(
        html.unescape(m.decode())
        for m in re.findall(rb"<w:delText(?:\s[^>]*)?>(.*?)</w:delText>", new)
    )
    if new_del > 0 and new_ins > 0 and word in deltexts:
        kind = "tracked"
    elif new_del == 0 and new_ins == 0:
        kind = "overwrite"
    else:
        kind = "other"
    other_text = [
        i
        for i, (a, b) in enumerate(zip(old_ps, new_ps))
        if i != idx and accepted_text(a) != accepted_text(b)
    ]
    return {
        "target_ok": text_ok and not other_text,
        "kind": kind,
        "new_w_ins": new_ins,
        "new_w_del": new_del,
        "target_text_ok": text_ok,
        "other_paragraph_text_changed": other_text,
    }


TEXT_LEAVES = {
    "t",
    "delText",
    "instrText",
    "delInstrText",
}  # whitespace inside these is content


def _normalise(root: etree._Element, unordered: bool) -> etree._Element:
    """Drop whitespace-only text/tails outside text leaves (indentation; insignificant in OOXML,
    which has no mixed content outside w:t-like leaves), and for [Content_Types].xml / *.rels sort
    the children (their order carries no meaning)."""
    for el in root.iter():
        if not isinstance(el.tag, str):
            continue
        if (
            el.text is not None
            and not el.text.strip()
            and etree.QName(el).localname not in TEXT_LEAVES
        ):
            el.text = None
        if el.tail is not None and not el.tail.strip():
            el.tail = None
    if unordered:
        kids = sorted(root, key=lambda k: etree.tostring(k, method="c14n"))
        for k in list(root):
            root.remove(k)
        root.extend(kids)
    return root


def _canon_part(name: str, b: bytes) -> bytes:
    """Canonical form of an XML part: inclusive C14N (drops the XML declaration and BOM, normalises
    attribute order, quote style and empty-element form) after _normalise(). Non-XML: raw bytes."""
    if not (name.endswith((".xml", ".rels"))):
        return b
    try:
        root = etree.fromstring(b)
    except etree.XMLSyntaxError:
        return b
    unordered = name == "[Content_Types].xml" or name.endswith(".rels")
    return etree.tostring(_normalise(root, unordered), method="c14n")


def _canon_paragraphs(xml: bytes) -> list[bytes]:
    """Canonical form of each outermost w:p: exclusive C14N (so root-level namespace declarations a
    paragraph inherits but doesn't use don't count) after _normalise()."""
    root = _normalise(etree.fromstring(xml), False)
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    out = []
    for p in root.iter(W + "p"):
        if any(a.tag == W + "p" for a in p.iterancestors()):
            continue
        out.append(etree.tostring(p, method="c14n", exclusive=True))
    return out


def canonical_noop(before: bytes, after: bytes) -> dict[str, Any]:
    a, b = parts(before), parts(after)
    same = sum(1 for n in a if n in b and _canon_part(n, a[n]) == _canon_part(n, b[n]))
    return {
        "parts_canonically_identical": same,
        "parts_total": len(a),
        "canonically_changed": sorted(
            n for n in a if n not in b or _canon_part(n, a[n]) != _canon_part(n, b[n])
        ),
        "parts_added": sorted(set(b) - set(a)),
    }


def canonical_edit(before: bytes, after: bytes) -> dict[str, Any]:
    a, b = parts(before), parts(after)
    other = sorted(
        n
        for n in set(a) | set(b)
        if n != MAIN
        and (n not in a or n not in b or _canon_part(n, a[n]) != _canon_part(n, b[n]))
    )
    pa, pb = _canon_paragraphs(a[MAIN]), _canon_paragraphs(b[MAIN])
    changed = [i for i, (x, y) in enumerate(zip(pa, pb)) if x != y]
    return {
        "other_entries_canonically_changed": other,
        "paragraphs_canonically_changed": changed if len(pa) == len(pb) else None,
        "paragraph_count_before": len(pa),
        "paragraph_count_after": len(pb),
    }


def structural_valid(out: bytes) -> dict[str, Any]:
    r: dict[str, Any] = {}
    try:
        with zipfile.ZipFile(io.BytesIO(out)) as z:
            r["zip_ok"] = z.testzip() is None
            bad = []
            for n in z.namelist():
                if n.endswith((".xml", ".rels")):
                    try:
                        etree.fromstring(z.read(n))
                    except etree.XMLSyntaxError as e:
                        bad.append(f"{n}: {e}")
            r["xml_parse_errors"] = bad
    except zipfile.BadZipFile as e:
        r["zip_ok"] = False
        r["xml_parse_errors"] = [str(e)]
    try:
        docx.Document(io.BytesIO(out))
        r["python_docx_loads"] = True
    except Exception as e:  # noqa: BLE001
        r["python_docx_loads"] = False
        r["python_docx_error"] = f"{type(e).__name__}: {e}"[:300]
    return r


def ooxml_validate(out: bytes, baseline: bytes) -> dict[str, Any]:
    """Open XML SDK validation via Docxodus verify_deliverable; counts only errors the output
    introduced (disposition != preExisting), in the validity categories, as tools/oracle does."""
    try:
        res = scalpel_call(
            "verify_deliverable",
            {
                "docxB64": base64.b64encode(out).decode(),
                "baselineB64": base64.b64encode(baseline).decode(),
            },
        )
    except Exception as e:  # noqa: BLE001
        return {"ok": None, "exception": f"{type(e).__name__}: {e}"[:300]}
    findings = res.get("findings", [])
    errs = [
        f
        for f in findings
        if (f.get("severity") == "error" or f.get("blocksDelivery"))
        and f.get("disposition") not in ("preExisting", "resolved")
        and f.get("category") in VALIDITY_CATEGORIES
    ]
    keep = ("code", "category", "disposition", "message", "owningPartUri")
    return {
        "ok": not errs,
        "decision": res.get("decision"),
        "new_errors": [{k: f.get(k) for k in keep if f.get(k)} for f in errs][:10],
        "new_error_count": len(errs),
    }


# --------------------------------------------------------------------------- driver


def run(fn: Callable[..., bytes], *args: Any) -> tuple[bytes | None, dict[str, Any]]:
    t0 = time.perf_counter()
    try:
        out = fn(*args)
        return out, {"ok": True, "seconds": round(time.perf_counter() - t0, 3)}
    except Exception as e:  # noqa: BLE001
        return None, {
            "ok": False,
            "seconds": round(time.perf_counter() - t0, 3),
            "error": f"{type(e).__name__}: {e}"[:500],
            "trace": traceback.format_exc()[-1500:],
        }


def bench_file(path: Path, validate: bool) -> dict[str, Any]:
    data = path.read_bytes()
    row: dict[str, Any] = {"file": str(path.relative_to(ROOT)), "bytes": len(data)}
    word = pick_word(path)
    row["word"] = word
    idx = target_index(path, word) if word else None
    row["target_paragraph"] = idx
    row["engines"] = {}
    row["tokens"] = {
        "document_xml": len(ENC.encode(parts(data)[MAIN].decode("utf-8", "replace")))
    }
    for eng in ENGINES:
        e: dict[str, Any] = {}
        out, e["noop_run"] = run(NOOP[eng], data)
        if out is not None:
            e["noop"] = noop(data, out)
            e["noop_canonical"] = canonical_noop(data, out)
            e["noop_valid"] = structural_valid(out)
            if validate and out != data:
                e["noop_valid"]["ooxml"] = ooxml_validate(out, data)
            elif validate:
                e["noop_valid"]["ooxml"] = {
                    "ok": True,
                    "note": "byte-identical to input",
                }
        if word:
            out, e["edit_run"] = run(EDIT[eng], data, word)
            if out is not None:
                e["edit"] = edit_locality(data, out)
                e["edit"].update(classify_edit(data, out, idx, word))
                e["edit_canonical"] = canonical_edit(data, out)
                e["edit_valid"] = structural_valid(out)
                if validate:
                    e["edit_valid"]["ooxml"] = ooxml_validate(out, data)
        try:
            row["tokens"][eng] = len(ENC.encode(read_view(eng, data)))
        except Exception as ex:  # noqa: BLE001
            row["tokens"][eng] = None
            e["read_error"] = f"{type(ex).__name__}: {ex}"[:300]
        row["engines"][eng] = e
    print(f"done {path.name}", file=sys.stderr, flush=True)
    return row


def versions() -> dict[str, str]:
    from importlib.metadata import version

    return {
        "quillwork": quillwork.__version__,
        "adeu": adeu.__version__,
        "docx-scalpel (Docxodus client)": version("docx-scalpel"),
        "python-docx": version("python-docx"),
        "lxml": version("lxml"),
        "tiktoken": tiktoken.__version__,
        "python": platform.python_version(),
        "os": f"{platform.system()} {platform.release()} {platform.machine()} (macOS {platform.mac_ver()[0]})",
        "date": time.strftime("%Y-%m-%d"),
    }


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("files", nargs="*", type=Path)
    ap.add_argument("--json", type=Path, default=HERE / "results-compare.json")
    ap.add_argument("--no-validate", action="store_true")
    args = ap.parse_args(argv)
    files = args.files or sorted(
        (ROOT / "corpus" / "synthetic").glob("*.docx")
    ) + sorted((ROOT / "corpus" / "fixtures").glob("*.docx"))
    rows = [bench_file(p.resolve(), not args.no_validate) for p in files]
    out = {"env": versions(), "edit_mode": EDIT_MODE, "rows": rows}
    args.json.write_text(json.dumps(out, indent=1, ensure_ascii=False) + "\n")
    print(f"wrote {args.json}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
